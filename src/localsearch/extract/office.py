r"""Office 与 PDF 四类容器（`F3-02`）：pdf / docx / xlsx / pptx → 正文。

四类容器，四个库（决策 D6）：PDF 走 PyMuPDF，`.docx` 走 python-docx，`.xlsx` 走
openpyxl，`.pptx` 走 python-pptx。共同点只有三条：**字节已经读进来了**（走
`base.read_bytes()`，`E27` 那条 `\\?\` 的债只在一处接）、**正文逐行攒进 `_Text`**
（上限与截断见下）、**失败一律翻译成 `ParseFailed`**（带上是哪个库、坏在哪一步）。

------------------------------------------------------------------------------
收哪四类，`.doc` / `.xls` / `.ppt` 为什么不收
------------------------------------------------------------------------------

老式二进制 Office（`.doc` / `.xls` / `.ppt`）**不在名单里**，而且不是"登记了但打不开"，
是**根本不认领**：它们会落到 `UnsupportedFormat`，`P4` 据此把它们留在 `NAME_ONLY`
（文件名照搜），**不进失败报表**。

理由在 `base.ParseFailed` 的文档里那张表：这几个不是坏文件，是 v1 不做的格式。认领了
再抛失败，等于把"我们没做"记成"用户的文件坏了"——报表里的噪音是这么来的。真要读它们，
要么加 `olefile` 那一套，要么走本机 Office/WPS 转换，都是 `P9` 之后的事。

⚠ 但 `.docx` 与 `.doc` 只差一个字母，用户分不清：**搜不到 `.doc` 里的字是预期行为，
不是 bug**。这条写进 `PROGRESS.md` 的遗留里。

------------------------------------------------------------------------------
容器标记：正文里为什么会有【…】
------------------------------------------------------------------------------

正文里会出现这几种**我们自己加的**标记行：

| 标记 | 什么时候加 |
|---|---|
| `【第 N 页】` | PDF，每页正文之前 |
| `【工作表：名】` | xlsx，每个工作表之前 |
| `【第 N 张幻灯片】` | pptx，每张幻灯片之前 |
| `【第 N 张幻灯片·备注】` | pptx，该页的演讲者备注之前（备注非空才加） |

三个理由：

1. **不标记的话这些结构在下游就没了。** 切片只带 `char_start` / `char_end`，那是
   `ExtractedDoc.text` 上的下标，**回不到 PDF 的页码**（一份 PDF 的"第 7 页"在正文里
   只是个偏移）。等 `P7` 要显示"命中在第 7 页"时，唯一的来源就是正文里这句话。
2. **它不是伪造内容。** 页号来自 PDF 自己的页对象、表名来自工作簿自己的
   `sheetnames`、张数来自幻灯片的顺序——都是文件里**真有**的事实，只是文件本身没把它
   写进文字流。加的是"这里换了一页"，不是"这里有一句话"。
3. **行为统一**：四类里只有 docx 不加，理由是 docx **没有页这个概念**——页码是 Word
   排版时才算出来的，文件里不存在。给它编一个页码就是真伪造了。

标记行和正文一样**参与偏移与切片**（`Chunk.char_start` 指向它也算数），所以"跳到原文"
能用；代价是切片文本里可能混进标记，与 `plaintext.py` 不清理行尾那条取舍同类。

------------------------------------------------------------------------------
上限与截断：`MAX_CHARS`
------------------------------------------------------------------------------

`_Text` 攒到 `MAX_CHARS` 就**整行地停**（不在行内切），此时 `ExtractedDoc.truncated`
为真、并留一条警告。这类在 `F3-03` 之后搬到了 `extract/_text.py`——`web_mail.py`
要的是同一套截断语义，而"整行地停"与 `has_body` 的判定抄成两份迟早分岔。

这么做的理由是**容器的解压放大**：`[scan] max_file_mb` 限的是
**容器本身的字节数**，而一个 64 MB 的 xlsx 解开可以是几 GB 的 XML。上限取 1e6 字符的
两个理由：

- **够用**：`D8` 的 200 片上限 × `target_tokens` 512 × (1 - overlap 0.15) ≈ 8.7 万 token
  会被真正切到。按 ASCII 约 4 字符/token 算，最坏也只有约 35 万字符——1e6 是它的
  2.9 倍，**切不到的东西不会因为上限而丢**。
- **内存**：1e6 字符的 `str` 在 CPython 里是 2–4 MB，可以忽略。

⚠ **这个上限只在"正文累积"这一层**，它挡不住的两件事要明说：① `docx` / `pptx` / `pdf`
的库会把整棵文档树先读进内存，压缩放大那一头**挡不住**（xlsx 走的是流式
`read_only=True`，那一头挡得住）；② 它不改变"这个文件有多大"——`bytes` 记的是原始
字节数。v1 接受，`zip` 炸弹是另一个层面的问题（记在遗留里）。

------------------------------------------------------------------------------
解析成功但正文是空的：留一条警告
------------------------------------------------------------------------------

四类容器都会：正文（去掉空白）为空时**追加一条 `warnings`**。这条与 `plaintext.py`
的口径不同——那里 0 字节文件是**静默的成功**（"空文件"是它自己的事实）。这里不同：
一个 40 MB 的 PDF 一页页全是扫描图，正文当然是空的，而**用户以为全文索引是完整的**。
两者的区别不是洁癖，是"用户会不会有一个错误的期待"。

v1 不做 OCR（计划 §1.3 划在 P8），所以这条警告里直接说清楚"图片里的字搜不到"。

------------------------------------------------------------------------------
库按需导入，不在模块顶层
------------------------------------------------------------------------------

四个库都在各自的 `_parse_*` **函数里**导入，不是模块顶层。理由是量过的
（`tools/bench_imports.py`，**已入库**：每个候选项起一个**新进程**、量"从 import 到
import 完成"，每项三轮取最小；各轮读数、内存占用率与两次运行都在 §3 的 **E32**）：

| 导入 | 最小 | 三次读数 |
|---|---|---|
| `import localsearch.extract.plaintext`（基线） | 90.3 ms | 91.1 / 98.6 / 90.3 |
| `pymupdf` | 168.5 ms | 174.7 / 182.2 / 168.5 |
| `docx` | 145.0 ms | 156.3 / 154.5 / 145.0 |
| `openpyxl` | 335.8 ms | 335.8 / 340.8 / 346.1 |
| `pptx` | 249.3 ms | 255.0 / 259.3 / 249.3 |
| `lxml.etree` | 58.2 ms | 63.2 / 61.6 / 58.2 |
| **四个一起** | **591.0 ms** | 628.3 / 620.2 / 591.0 |

`extract/__init__.py` 的 `build_registry()` 会 import 本模块，所以在顶层写 import
等于让**每一次 `lsearch` 启动**（包括 `status` / `doctor` / 一次检索）都多付这约
620 ms——**基线的 6 倍多**，而 `plaintext` 才是每次都要用的那一个。挪进函数之后，
只有真的提取到这类文件时才付。

（这张表原先是"一次性内联命令量过"的一组数——`191.9 / 137.2 / 360.0 / 256.0 /
64.6 / 599.2`，与基线 `91.9`。`2026-10-15` 补了可复跑的脚本之后按 `E12` 换了源：
量级与结论一字未变，但复跑得出来了。**三次里 `四个一起` 是 628.3 / 620.2 / 591.0**，
极差 6.3%，属 `E17` 那类方差——不要拿某一次的数当准值。）

代价是 `except` 子句里的异常类也要在函数里取——那个库抛什么异常本来就只有它自己
知道，写在函数里反而离事实更近。
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from io import BytesIO
from itertools import repeat
from typing import Any, Callable, Iterable

from localsearch.core.models import ExtractedDoc

from ._text import _Text
from .base import (
    ExtractConfig,
    Extractor,
    ParseFailed,
    UnsupportedFormat,
    ext_of,
    read_bytes,
)

#: 收的扩展名。**只有这四个**——`.doc` / `.xls` / `.ppt` 的理由见模块文档。
OFFICE_EXTS = frozenset({"pdf", "docx", "xlsx", "pptx"})

#: 正文的字符上限，理由（200 片的算术 + 内存）见模块文档。
MAX_CHARS = 1_000_000


# ==============================================================================
# WordprocessingML 的标签名：直接写 Clark 记法，不走 docx 的 qn()
# ==============================================================================

#: 命名空间固定（OOXML 规范的一部分），标签名也固定，所以这几个常量不需要依赖
#: `docx` 才能算出来——这正好让解析正文的那几个小函数**不必导入 docx**。
#: `tests/unit/test_extract_office.py` 里有一条断言拿 `qn()` 逐条核对它们，
#: 万一哪天写成别的名字，那条测试会红，而不是安静地把正文收成空。
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_P = f"{{{_W_NS}}}p"
_W_T = f"{{{_W_NS}}}t"
_W_TAB = f"{{{_W_NS}}}tab"
_W_BR = f"{{{_W_NS}}}br"
_W_CR = f"{{{_W_NS}}}cr"


# ==============================================================================
# 解析结果
# ==============================================================================


@dataclass(frozen=True, slots=True)
class _Parsed:
    """一个容器解析出来的东西。`bytes` / `chars` / `format` 由 `extract()` 统一补。"""

    text: str
    title: str | None
    meta: dict[str, Any]
    warnings: tuple[str, ...] = ()


# ==============================================================================
# 小工具
# ==============================================================================


def _clean_title(value: Any) -> str | None:
    """收敛成标题：去空白、空串算没有。

    **去换行**是必须的：这个值会进嵌入前缀的 `title: {t} | text: ...`（`M3`），
    标题里的换行会与那个格式的分隔符撞上。所以这里用 `" ".join(split())` 把内部
    空白也压成一个空格。
    """
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _cell_text(value: Any) -> str:
    """表格单元格的字符串化。`str()` 就够——数字格式（千分位、日期显示式样）是 xlsx
    的**显示层**，v1 不复刻（记在遗留里）。内部空白压平（多行的格子也该只占一行）。"""
    return " ".join(str(value).split())


#: 见 `_xlsx_text`：无限长的一行 `None`，用来配"这一遍没有第二遍可比"。
_NO_ROW: Any = repeat(None)


def _row_line(values: Any) -> str:
    """一行格子拼成一行正文：**空的格子直接丢掉**，用 `\\t` 接剩下的。

    不拿制表符给空格子占位——正文是给人搜的，不是给人对齐的，而连续的 `\\t\\t`
    除了让 FTS 多几条空 token 没有任何用。
    """
    cells = [_cell_text(v) for v in values]
    return "\t".join(c for c in cells if c)


def _inside_paragraph(node: Any) -> bool:
    """这个 `w:p` 是不是套在另一个 `w:p` 里。

    文本框（`w:pict > v:shape > v:textbox > w:txbxContent > w:p`）、以及别的宿主
    结构里的段落都会这样**嵌套**。外层段落用 `iter()` 递归取 `w:t` 时会把里层的字
    一并取走，所以里层那一遍必须跳过——不然同一句话在正文里出现两次。
    """
    parent = node.getparent()
    while parent is not None:
        if parent.tag == _W_P:
            return True
        parent = parent.getparent()
    return False


def _paragraph_text(node: Any) -> str:
    r"""一个 `w:p` 的文本。

    走 `iter()` 是**先位置后内容**的取法（文档顺序），所以文本框里的字出现在图形
    锚定的那个位置，而不是全文最后。顺带认三种行内元素：`w:tab` → `\t`、`w:br` /
    `w:cr` → `\n`。**不认** `w:delText`（修订里被删掉的字）与 `w:instrText`（域代码）
    ——它们是另两种标签，`iter()` 取的是 `w:t`，天然不含。

    ⚠ 表格里的段落也在 `body.iter(w:p)` 的结果里（每个单元格一段），所以表格文字
    与正文**按文档顺序混排**，不额外标记"这里进表格了"。这是有意的：v1 的正文是
    给检索用的，不是给人复排版面的。
    """
    parts: list[str] = []
    for element in node.iter():
        tag = element.tag
        if tag == _W_T:
            parts.append(element.text or "")
        elif tag == _W_TAB:
            parts.append("\t")
        elif tag == _W_BR or tag == _W_CR:
            parts.append("\n")
    return "".join(parts)


def _page_marker(number: int) -> str:
    return f"【第 {number} 页】"


def _sheet_marker(name: Any) -> str:
    return f"【工作表：{' '.join(str(name).split())}】"


def _slide_marker(number: int) -> str:
    return f"【第 {number} 张幻灯片】"


def _notes_marker(number: int) -> str:
    return f"【第 {number} 张幻灯片·备注】"


def _pdf_errors() -> tuple[type[BaseException], ...]:
    """PyMuPDF 抛错有**两族**，共同祖先只到 `Exception` / `RuntimeError`，所以两族都要列：

    - `pymupdf.FileDataError`（基类是 `RuntimeError`）：打不开、结构坏了、空文件
      （`EmptyFileError` 是它的子类）；
    - `pymupdf.mupdf.FzErrorBase`（基类是 `Exception`）：MuPDF 自己的错，取文本时
      才可能抛出来（`FzErrorFormat` 等）。
    """
    import pymupdf

    return (pymupdf.FileDataError, pymupdf.mupdf.FzErrorBase)


def _ooxml_errors() -> tuple[type[BaseException], ...]:
    """OOXML 坏掉会抛的那几种，实测过的四种形态（`2026-10-15`）：

    | 文件长什么样 | 谁抛的 |
    |---|---|
    | 根本不是 zip（垃圾字节、截断的文件） | `zipfile.BadZipFile` |
    | 是 zip，但没有 `[Content_Types].xml` | `KeyError`（消息是库自己的） |
    | `document.xml` / `slide1.xml` 本身是坏的 XML | `lxml.etree.XMLSyntaxError` |
    | `sheet1.xml` 本身是坏的 XML | `xml.etree.ElementTree.ParseError` |

    后两种分开列是**实测出来的分工**：python-docx 与 python-pptx 走 lxml、openpyxl
    走 stdlib 的 `ElementTree`，而 `lxml` 的 `XMLSyntaxError` **不是** `ET.ParseError`
    的子类（它的 mro 里另有一条 `SyntaxError`）。少列一个的后果是那个形态的坏文件变成
    未捕获异常、把整轮索引带下去（`A6` 要的正好相反）。

    库抛的其它异常**继续往上冒**（`base` 那条"真接不住的不要接"）：这四种是"这个文件
    坏了"的确定信号，不是"什么都接住"的网。
    """
    import xml.etree.ElementTree as ET

    from lxml.etree import LxmlError

    return (zipfile.BadZipFile, KeyError, ET.ParseError, LxmlError)


def _failed(kind: str, exc: BaseException) -> ParseFailed:
    """把库的异常翻成一条人能读的理由。**带上库自己说的话**——那是唯一的线索。"""
    return ParseFailed(f"{kind} 解析失败（{type(exc).__name__}：{exc}）")


# ==============================================================================
# 四种容器
# ==============================================================================


def _parse_pdf(data: bytes, cap: _Text) -> _Parsed:
    """PDF：逐页取文本层，每页一行标记。

    `page.get_text()` 用的是 PyMuPDF 的默认模式（按块，不重排）。**不做**版面分析
    （分栏文档的阅读顺序会乱，`get_text(sort=True)` 也治不了根本），也**不取注释、
    不取表单域**——v1 只要文字层。
    """
    import pymupdf

    empty_pages = 0
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ParseFailed(
                    "PDF 有密码 / 被加密，v1 不解密（`Doc.needs_pass` 为真）。"
                    "先在本机用什么打开一次、去掉密码再索引"
                )
            pages = doc.page_count
            if pages == 0:
                # 实测：`b"%PDF-1.4\n...trailer\n<<>>\n%%EOF\n"` 这种**头对但内容空**的
                # 文件 MuPDF 会照收，页数是 0（`2026-10-15`）。放它过去的话，下面的
                # "每一页都没有文本层"那条警告会对着**不存在的页**说话——那是假话。
                raise ParseFailed(
                    "PDF 解析出来一页都没有（结构坏了，或被解析器当空壳收下了）"
                )
            info = doc.metadata or {}
            for number, page in enumerate(doc, start=1):
                body = page.get_text()
                if body.strip():
                    cap.marker(_page_marker(number))
                    cap.block(body)
                else:
                    empty_pages += 1
                if cap.full:
                    break
    except _pdf_errors() as exc:
        raise _failed("PDF", exc) from exc

    meta: dict[str, Any] = {"pages": pages, "empty_pages": empty_pages}
    for key in ("producer", "creator"):
        value = _clean_title(info.get(key))
        if value:
            meta[key] = value

    warnings: list[str] = []
    # 截断时不再报"哪些页是空的"——那些没走到的页不是空的，只是没看
    if empty_pages and not cap.truncated:
        if empty_pages < pages:
            warnings.append(
                f"{empty_pages} / {pages} 页没有文本层（扫描件或纯图片页），"
                "这些页的正文缺失"
            )
    return _Parsed(cap.text, _clean_title(info.get("title")), meta, tuple(warnings))


def _parse_docx(data: bytes, cap: _Text) -> _Parsed:
    r"""docx：按文档顺序取每个顶层的 `w:p`，一段一行（表格里的段落也在内）。

    走 `body.iter(_W_P)` 而不是 `Document.paragraphs`，是因为后者**只给正文区**：
    表格里的字、文本框里的字都不在里面，而它们是真内容（`2026-10-15` 实测：
    `paragraphs` 给 3 段，`body.iter` 给表格与文本框的字）。

    **不取**页眉 / 页脚（`section.header` / `footer`）：它们常常每页重复（公司名、
    页码），重复文本灌进 FTS 只会把评分拉歪。**不取**批注与修订记录。

    `meta["paragraphs"]` 是**收进来的 `w:p` 个数**，所以表格单元格里的段落也各算一个
    （它们本来就是各自成段的）——这个数不是"文档有几个自然段"，是"正文由几行拼成"。
    """
    from docx import Document

    try:
        doc = Document(BytesIO(data))
        paragraphs = 0
        for node in doc.element.body.iter(_W_P):
            if _inside_paragraph(node):
                continue
            cap.line(_paragraph_text(node))
            paragraphs += 1
            if cap.full:
                break
        title = _clean_title(doc.core_properties.title)
        tables = len(doc.tables)
    except _ooxml_errors() as exc:
        raise _failed("docx", exc) from exc

    return _Parsed(cap.text, title, {"paragraphs": paragraphs, "tables": tables})


def _xlsx_text(wb: Any, cap: _Text, mirror: Any = None) -> tuple[int, int]:
    """把工作簿摊成正文，返回 `(公式格数, 缓存缺失格数)`。

    两遍读的落地（见 `_parse_xlsx`）：`mirror` 为 `None` 时是第一遍，数的是"有多少格
    是公式"；`mirror` 给了第一遍的工作簿时是第二遍，数的是"有多少格的缓存缺失、退回
    了公式原文"。两遍用 `zip()` **逐行对齐**（两个都是流式迭代器），所以不需要把
    格子坐标存下来——那才是会吃内存的写法。

    ⚠ 这里有一条**载重假设**：同一个 `read_only` 工作表**可以被重新迭代**（第二遍要
    在 `mirror` 上从头再走一遍）。`2026-10-15` 实测过 openpyxl 3.1.5 的行为：同一个
    工作表调三次 `iter_rows()` 三次都给全量数据，且**被打断的那个迭代器**（只取了一行）
    仍能继续给出剩下的。这条要是哪天变了，症状是第二遍的正文**只剩标记行**——所以
    `test_a_formula_cell_with_a_cached_value_uses_the_value` 那条测试是它的回归线
    （它要的值来自第二遍，第二遍空了它就红）。
    """
    formulas = 0
    missing = 0
    for name in wb.sheetnames:
        cap.marker(_sheet_marker(name))
        if cap.full:
            break
        rows = wb[name].iter_rows(values_only=True)
        # `_NO_ROW` 是"无限长的一行 `None`"：`zip(row, 它)` 让每个格子都配到 `None`，
        # 等价于"这一遍没有第一遍可比"（用 `repeat(None)` 是错的——那配出来的是
        # 一个 `None` 行、`zip` 解不开，`2026-10-15` 冒烟时踩到过）。
        # `itertools.repeat` 无限且不消耗，所以每一行共用同一个对象是安全的。
        others: Iterable[Iterable[Any]] = (
            repeat(_NO_ROW) if mirror is None else mirror[name].iter_rows(values_only=True)
        )
        for row, other in zip(rows, others):
            cells: list[str] = []
            for value, original in zip(row, other):
                if value is None and original is not None:
                    # 缓存缺失（这个文件不是 Excel 保存的？）→ 退回公式原文，
                    # 宁可搜到 `=SUM(B2:B2)` 也不要**静默**丢掉这一格
                    value = original
                    missing += 1
                if mirror is None and isinstance(value, str) and value.startswith("="):
                    formulas += 1        # 只有第一遍数它；第二遍见到的公式都是"退回来的"
                if value is None or value == "":
                    continue
                text = _cell_text(value)
                if text:
                    cells.append(text)
            if cells:
                cap.line("\t".join(cells))
            if cap.full:
                break
    return formulas, missing


def _parse_xlsx(data: bytes, cap: _Text) -> _Parsed:
    """xlsx：每个工作表一段，一行格子一行正文。

    公式这一档**读两遍**：第一遍 `data_only=False`（正文里是 `=SUM(...)`），顺便
    回答"这个工作簿里有没有公式"；有公式才读第二遍 `data_only=True`，**优先取缓存值**
    （Excel 保存时算好的结果），缓存缺失的格子退回公式原文并计数。

    为什么非要两遍：单用 `data_only=True`，那些没有缓存的格子（openpyxl / pandas 写
    的文件就是这样）会**静默变成空格子**——一整列算出来的数就没了，而搜不到不会有
    任何提示；单用 `data_only=False`，则"搜合计那一栏的数字"永远搜不到，而那恰恰是
    表格最常见的用法。代价写在这里：**有公式的工作簿要多读一遍**（没有公式时第一遍
    的正文直接用，不多花）；第一遍就已经攒满上限时也不读第二遍——正文已经满了，
    第二遍的值一个也进不来。

    隐藏表**照收**（`sheet_state == "hidden"`）：数据在隐藏表里是常态，"隐藏"是排版
    决定，不是"这些内容不算数"。
    """
    from openpyxl import load_workbook

    try:
        plain = load_workbook(BytesIO(data), read_only=True, data_only=False)
        title = _clean_title(plain.properties.title if plain.properties else None)
        formulas, _ = _xlsx_text(plain, cap)
        missing = 0
        if formulas and not cap.full:
            cached = load_workbook(BytesIO(data), read_only=True, data_only=True)
            cap.reset()                    # 见 `_Text.reset()`：必须是重置，不是换对象
            _, missing = _xlsx_text(cached, cap, mirror=plain)
        sheets = len(plain.sheetnames)
    except _ooxml_errors() as exc:
        raise _failed("xlsx", exc) from exc

    warnings: list[str] = []
    if missing:
        warnings.append(
            f"{missing} 个格子是公式但文件里没有算好的值（这个文件不是 Excel 存的？），"
            "这些格子用的是公式原文"
        )
    return _Parsed(
        cap.text, title, {"sheets": sheets, "formulas": formulas}, tuple(warnings)
    )


def _parse_pptx(data: bytes, cap: _Text) -> _Parsed:
    """pptx：每张幻灯片一段，形状按文件里的顺序取（组合形状递归进去）。

    取三种形状的文本：文本框 / 占位符（`has_text_frame`）、表格（逐行、格子用 `\\t`
    接）、组合（`GROUP`，递归）。图片与图表**不取**（v1 不做 OCR，图表的数据在
    xlsx 里另有正文）。演讲者备注**取**——那是用户真的写过、也真的会去搜的字。
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(BytesIO(data))
        slides = len(prs.slides)
        notes = 0
        for number, slide in enumerate(prs.slides, start=1):
            cap.marker(_slide_marker(number))
            _ppt_shapes(slide.shapes, cap, MSO_SHAPE_TYPE)
            if slide.has_notes_slide:
                text = slide.notes_slide.notes_text_frame.text.strip()
                if text:
                    cap.marker(_notes_marker(number))
                    cap.block(text)
                    notes += 1
            if cap.full:
                break
    except _ooxml_errors() as exc:
        raise _failed("pptx", exc) from exc

    return _Parsed(cap.text, _clean_title(prs.core_properties.title), {"slides": slides, "notes": notes})


def _ppt_shapes(shapes: Any, cap: _Text, group: Any) -> None:
    """一层形状。组合形状递归进去（顺序就是文件里的顺序）。"""
    for shape in shapes:
        if shape.shape_type == group.GROUP:
            _ppt_shapes(shape.shapes, cap, group)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                line = _row_line([cell.text for cell in row.cells])
                if line:
                    cap.line(line)
        elif shape.has_text_frame:
            cap.block(shape.text_frame.text)
        if cap.full:
            return


#: 四类的入口。**键就是认领的扩展名**，`supports()` 与它共用一份名单（`OFFICE_EXTS`）。
_PARSERS: dict[str, Callable[[bytes, _Text], _Parsed]] = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "pptx": _parse_pptx,
}

#: 正文为空时那句警告。分成四句是因为"为什么是空的"**因格式而异**，而这句话就是
#: 用户排查时唯一能看见的线索。
_EMPTY_HINTS = {
    "pdf": "每一页都没有文本层（扫描件？）；v1 不做 OCR，图片里的字搜不到",
    "docx": "文档里没有文字（只有图片？）；v1 不做 OCR",
    "xlsx": "工作簿里没有单元格文字（只有图片或图表？）",
    "pptx": "幻灯片里没有文字（只有图片？）；v1 不做 OCR",
}


class OfficeExtractor:
    """四类容器 → 正文。四个库各管一段，共同的三条见模块文档。"""

    name = "office"

    def __init__(self, config: ExtractConfig, *, max_chars: int = MAX_CHARS) -> None:
        self._config = config
        #: 可注入，默认 `MAX_CHARS`：测试要能在**不造一个 1 MB 的文件**的前提下把那条
        #: 截断路径走一遍。（与 `chunk` 那个 token 计数器同一个写法：可注入 + 默认值。）
        self._max_chars = max_chars

    def supports(self, ext: str, mime: str) -> bool:
        """只有那四个扩展名。`mime` 一个字符不用（理由在 `base.mime_of()`）。"""
        del mime
        return ext in OFFICE_EXTS

    def extract(self, path: str) -> ExtractedDoc:
        ext = ext_of(path)
        parse = _PARSERS.get(ext)
        if parse is None:
            # 走到这里说明有人绕过了 `supports()`（`extract_file()` 不会）。照实说，
            # 不假装成功——"注册了但返回空正文"是最难查的一种失败。
            raise UnsupportedFormat(f"office 不收 .{ext}（supports() 本该拦住它）")

        data = read_bytes(path, config=self._config)
        cap = _Text(self._max_chars)
        parsed = parse(data, cap)

        warnings = list(parsed.warnings)
        if cap.truncated:
            warnings.append(
                f"正文超过 {self._max_chars} 字符的上限，后面的内容没有提取"
                "（上限是 office.py 的 MAX_CHARS）"
            )
        if not cap.has_body:
            warnings.append(f"解析成功但没有提取到任何文本：{_EMPTY_HINTS[ext]}")

        return ExtractedDoc(
            text=parsed.text,
            title=parsed.title,
            meta={**parsed.meta, "format": ext, "bytes": len(data), "chars": len(parsed.text)},
            modality="text",
            source=self.name,
            truncated=cap.truncated,
            warnings=tuple(warnings),
        )


def build(config: ExtractConfig) -> Extractor | None:
    """按 `[extract] office` 决定建不建；关掉时返回 `None`（不注册，理由见 `plaintext.build()`）。"""
    if not config.office:
        return None
    return OfficeExtractor(config)
