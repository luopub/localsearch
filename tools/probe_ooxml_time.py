r"""三个 OOXML 库把"现在"写在哪：`F3-06` 的样本能不能逐字节重造，全看这一条。

```
PYTHONIOENCODING=utf-8 uv run python tools/probe_ooxml_time.py
```

------------------------------------------------------------------------------
为什么要有这个探针
------------------------------------------------------------------------------

`tests/fixtures/` 里有三份 OOXML 样本（`.docx` / `.xlsx` / `.pptx`），而 fixtures
的判据是"**提交上去的字节 = 生成器现在造出来的字节**"（`tools/make_fixtures.py
--check`）。所以"这三个库会不会写当前时间"不是个学术问题：写进去的话，同一份输入
两次生成的字节就不同，那份判据当场失效。

------------------------------------------------------------------------------
⚠ 这件事量错过一次，错在这里
------------------------------------------------------------------------------

第一版量法是"在**同一个进程**里连着造两份、比字节"，结论是"三份都稳定"。
**那个结论是假的**：两个格式的 `save()` 落在**同一秒**里，而它们写的是
`time.localtime()` 的**秒**级时间戳——同一秒当然一模一样。要看见它，两次生成
必须**跨过一秒的边界**（本脚本 `time.sleep(1.5)` 就是为这个加的），而且得在
**不同的进程**里造（同一进程里库会复用一些状态）。

这与 `E31` / `E34` 是同一条教训：**测的东西不是想测的东西，而它照样给出一个
看起来很确定的数**——所以探针里那个"睡 1.5 秒"不是保险，它是这个结论的全部。

------------------------------------------------------------------------------
它量什么
------------------------------------------------------------------------------

每个格式**起一个新进程**造一份，主进程比两次的三样东西：

1. 字节本身（以及长度与摘要——注意 `.xlsx` 那两行长度会差 1 字节：时间戳串
   `2026-10-15T20:42:03Z` 与 `…:04Z` 一样长，但 `created` 与 `modified` 都变了）；
2. **每个 zip 条目**的时间戳（`ZipInfo.date_time`）；
3. `docProps/core.xml` 的正文——里面 `dcterms:created` / `dcterms:modified`
   两个元素的值，脚本会把两版各打印一次。

结论（2026-10-16 04:41:36 与 04:42:00 两轮实测）：

| 格式 | zip 条目时间戳 | `docProps/core.xml` |
|---|---|---|
| `.docx` | **写"现在"** | 来自它自己的模板（`2013-12-23T23:15:00Z`），稳定 |
| `.xlsx` | **写"现在"** | **写"现在"**，`created` 与 `modified` 都是 |
| `.pptx` | **写"现在"** | 来自它自己的模板（`2013-01-27T09:14:16Z` / `…:09:15:58Z`），稳定 |

所以"三个都靠库自己稳定"是不成立的，**三份都得过一遍 `_normalize_zip()`**（它把
条目时间戳钉成 `1980-01-01`，并把 `core.xml` 里那两个元素的值钉成固定串）。

⚠ `.xlsx` 那一条还有一层：`openpyxl` 在 `save()` 里**无条件**覆写
`dcterms:modified`，**先给 `wb.properties.modified` 赋一个固定值也拦不住它**
（实测，见 `PROGRESS.md` §3 的 `E36`）——所以那一处只能在**打包之后**改，
没法在库这一层解决。

------------------------------------------------------------------------------
它不是什么
------------------------------------------------------------------------------

它不验"我们的样本是稳定的"——那件事的判据是 `make_fixtures.py --check`（0 漂移）
与 `tests/unit/test_fixtures.py`。这里量的是**库的性质**：哪一处会飘、飘在哪。
库换了版本之后重跑它，那张表就要核对一遍。

------------------------------------------------------------------------------
⚠ 打印路径上不许出现 `⚠`（这一条是本脚本自己踩出来的）
------------------------------------------------------------------------------

本脚本第一版在**打印**的那一行里带了 `U+26A0`（就是文档串里这个符号本身），于是
在**本机默认控制台**（代码页 936）下当场 `UnicodeEncodeError`、退出码 1——而
用 `PYTHONIOENCODING=utf-8` 跑的那几次完全正常（`E16` 记的就是这条，2026-10-15
`tools/scan_report.py` 踩过一次，症状是整轮输出全丢）。所以：**文档串里可以写，
`print` 的字面量里不行**，一路只用 GBK 编得了的字符。
"""

from __future__ import annotations

import datetime
import hashlib
import io
import subprocess
import sys
import time
import zipfile

from localsearch.core.console import use_safe_output

#: 子进程里造一份样本（只依赖一个参数：格式名），字节从 stdout 原样带回来。
#: 写成子进程而不是在进程内连着造两份，理由见模块文档的 ⚠。
CHILD = r'''
import io, sys

kind = sys.argv[1]
if kind == "docx":
    import docx
    document = docx.Document()
    document.add_paragraph("x")
    buffer = io.BytesIO()
    document.save(buffer)
elif kind == "xlsx":
    import openpyxl
    book = openpyxl.Workbook()
    book.active["A1"] = "x"
    buffer = io.BytesIO()
    book.save(buffer)
else:
    import pptx
    deck = pptx.Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    buffer = io.BytesIO()
    deck.save(buffer)
sys.stdout.buffer.write(buffer.getvalue())
'''

#: 两次生成之间要睡多久。**这个数是这个结论的全部**：不睡，两次落在同一秒，
#: 库写进去的秒级时间戳就成了"稳定"（第一版就是这么错的）。
GAP_SECONDS = 1.5

CORE = "docProps/core.xml"
TIME_TAGS = ("dcterms:created", "dcterms:modified")


def build(kind: str) -> bytes:
    """起一个**新进程**造一份样本，把字节带回来。"""
    result = subprocess.run([sys.executable, "-c", CHILD, kind], capture_output=True)
    if result.returncode != 0:
        raise SystemExit(result.stderr.decode("utf-8", "replace"))
    return result.stdout


def _core_times(blob: bytes) -> list[str]:
    """`core.xml` 里那两个元素的值，按出现顺序（没有就给空表）。"""
    text = blob.decode("utf-8", "replace")
    values: list[str] = []
    for tag in TIME_TAGS:
        at = text.find(f"<{tag}")
        if at < 0:
            continue
        start = text.find(">", at) + 1
        values.append(f"{tag} = {text[start:text.find('<', start)]}")
    return values


def _stamps(blob: bytes) -> list[tuple[int, ...]]:
    return [entry.date_time for entry in zipfile.ZipFile(io.BytesIO(blob)).infolist()]


def report(kind: str) -> None:
    first = build(kind)
    time.sleep(GAP_SECONDS)          # 跨过一秒的边界，见模块文档
    second = build(kind)

    first_zip = zipfile.ZipFile(io.BytesIO(first))
    second_zip = zipfile.ZipFile(io.BytesIO(second))
    core_first = first_zip.read(CORE) if CORE in first_zip.namelist() else b""
    core_second = second_zip.read(CORE) if CORE in second_zip.namelist() else b""

    print(f"== {kind}")
    print(f"   两次生成的字节相同：{first == second}"
          f"（{len(first)} / {len(second)} 字节，"
          f"{hashlib.sha256(first).hexdigest()[:12]} / "
          f"{hashlib.sha256(second).hexdigest()[:12]}）")
    stamps = _stamps(first)
    print(f"   zip 条目时间戳相同：{_stamps(first) == _stamps(second)}"
          f"（{len(stamps)} 个条目，第一份前 3 个 {stamps[:3]}）")
    print(f"   {CORE} 相同：{core_first == core_second}")
    for label, blob in (("第一份", core_first), ("第二份", core_second)):
        for line in _core_times(blob):
            print(f"      {label} {line}")


def main() -> int:
    # 输出面兜底（core/console.py）：管道 / 重定向下不许因多字节字符崩掉
    use_safe_output()
    print(f"机器时间（脚本自己取的）：{datetime.datetime.now():%Y-%m-%d %H:%M:%S}")
    print(f"两次生成之间睡 {GAP_SECONDS} 秒——不睡的话两次落在同一秒，看着「稳定」"
          "（见模块文档）")
    for kind in ("docx", "xlsx", "pptx"):
        report(kind)
    return 0


if __name__ == "__main__":  # pragma: no cover - 入口
    sys.exit(main())
