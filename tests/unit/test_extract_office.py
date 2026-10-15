r"""`extract/office.py`：四类容器的正文、标记、上限、失败（`F3-02`）。

------------------------------------------------------------------------------
这一组盯的六件事
------------------------------------------------------------------------------

1. **认领的边界**——四个扩展名认领，`.doc` / `.xls` / `.ppt` **不认**（它们走
   `UnsupportedFormat` → `NAME_ONLY`，理由在模块文档里），无扩展名也**不认**（那是
   `plaintext` 的）。最后两条比第一条重要：`office` 一旦认了 `.doc`，用户那些老式
   Word 文件会进失败报表；一旦认了无扩展名的文件，`README` 会被拿去当 PDF 解析。
2. **正文的组成**——标记行、文档顺序、空行、表格与文本框。这些是"正文长什么样"的
   全部约定，下游（切片、检索、`P7` 的页码跳转）都建立在它们上面。
3. **公式那一档**（xlsx 唯一有分歧的地方）——三种情况各一条：没公式、有公式但没缓存
   （退回公式原文 + 一条警告）、有公式且有缓存（**用缓存值**）。第三条同时是
   "同一个工作表可以重新迭代"那条载重假设的回归线：第二遍要是空了，它就红。
4. **上限与截断**——`max_chars` 注入一个小值，整行地停、`truncated` 为真、留一条警告；
   `[scan] max_file_mb` 那条上限也照旧管着 office。另外两条是 xlsx 两遍读专有的：
   两遍的长度可能**反过来**（缓存值比公式原文还长），`truncated` 必须跟着**最终**那份
   正文；以及第一遍就到顶时**不读第二遍**（这道闸门是 `_Text.reset()` 的前提）。
5. **失败一律是 `ParseFailed`**——五种坏法（不是 zip / 是 zip 但不是 OOXML / 坏
   `document.xml` / 坏 `slide1.xml` / 坏 `sheet1.xml`）各一条，消息里**必须带上是
   哪个库报的**；另有一条"头对但一页都没有的 PDF"（MuPDF 会照收，页数 0）。
6. **与库自己的表对齐**——`office.py` 里那几个手写的 WordprocessingML 标签常量必须
   与 `docx.oxml.ns.qn()` 给的值一致（写错的名字不会报错，只会让正文**安静地**变空）。

⚠ **这里造样本用的是被测库自己**（python-docx 写 docx、openpyxl 写 xlsx……），所以
"库写得对不对"这一层是自洽的、验不出来。独立样本（一旦落盘就不再变、跨库版本也能
复现）是 `F3-06` 的 fixtures 要补的东西，两者不重复：这一组验的是**语义**，那一组
验的是**字节**。

`tmp_path` 落在 `C:` 下——项目只允许碰 `C:` 与 `D:`。
"""

from __future__ import annotations

import io
import zipfile
from typing import Any, Iterable

import pytest

from localsearch.core.models import ExtractedDoc
from localsearch.extract import build_registry
from localsearch.extract.base import (
    ExtractConfig,
    ExtractError,
    FileUnreadable,
    ParseFailed,
    UnsupportedFormat,
    extract_file,
)
from localsearch.extract.office import (
    MAX_CHARS,
    OFFICE_EXTS,
    OfficeExtractor,
    _PARSERS,
    _W_BR,
    _W_CR,
    _W_NS,
    _W_P,
    _W_T,
    _W_TAB,
    build,
)

# ==============================================================================
# 造样本：用被测库自己写（这一层的自洽是已知的，见模块文档最后一段）
# ==============================================================================

#: PyMuPDF 自带的中文字体别名（简中）。不指定的话 `insert_text` 写不了中文，
#: 而这一组的语料就是中文。
CJK_FONT = "china-s"


def _pdf(
    path: Any,
    pages: Iterable[str | None],
    *,
    metadata: dict[str, str] | None = None,
) -> str:
    """一页一段文字；`None` / 空串 = 一页没有文字层（模拟扫描页）。"""
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text, fontname=CJK_FONT, fontsize=14)
    if metadata:
        doc.set_metadata(metadata)
    doc.save(str(path))
    doc.close()
    return str(path)


def _encrypted_pdf(path: Any, text: str = "加密页的正文") -> str:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text, fontname=CJK_FONT, fontsize=14)
    path.write_bytes(
        doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user"
        )
    )
    doc.close()
    return str(path)


def _docx(path: Any, document: Any) -> str:
    document.save(str(path))
    return str(path)


def _new_docx() -> Any:
    import docx

    return docx.Document()


def _xlsx(path: Any, workbook: Any) -> str:
    workbook.save(str(path))
    return str(path)


def _new_sheet(rows: Iterable[Iterable[Any]] = (), *, sheet: str = "台账") -> Any:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(list(row))
    return wb


def _new_pptx() -> Any:
    from pptx import Presentation

    return Presentation()


def _slide(prs: Any) -> Any:
    """加一张空白版式的幻灯片（版式 6 = 全空，自带占位符会往正文里加字）。"""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide: Any, text: str, *, top: int = 0) -> Any:
    """加一个写着 `text` 的文本框。注意别用 `add_slide` 自带的占位符——
    它们空着也会往正文里带提示文字（"单击此处添加标题"之类）。"""
    frame = slide.shapes.add_textbox(0, top, 300, 60).text_frame
    frame.text = text
    return frame


def _save_pptx(prs: Any, path: Any) -> str:
    prs.save(str(path))
    return str(path)


def _textbox_paragraph_xml(text: str) -> str:
    """一段"文本框"的 WordprocessingML：`w:pict > v:shape > v:textbox > w:txbxContent`。

    python-docx 造不出文本框（没有 API），所以只能塞一段原始 XML 进去——这也是
    这条测试唯一能验"文本框里的字收不收、会不会收两遍"的办法。
    """
    return (
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape style="width:100pt;height:50pt"><v:textbox><w:txbxContent>'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict>"
    )


def _add_textbox_paragraph(document: Any, text: str) -> None:
    from docx.oxml import parse_xml

    document.add_paragraph().add_run()._r.append(parse_xml(_textbox_paragraph_xml(text)))


#: 手写的 sheet XML：一格**公式 + 算好的缓存值**（`<f>` 与 `<v>` 并存）。
#: 只有真实的 Excel 存出来的文件才长这样，openpyxl 自己写的只有 `<f>` 没有 `<v>`
#: （所以测"优先用缓存值"这条规则必须自己造字节，不能用库写出来的文件）。
_SHEET_WITH_CACHED_FORMULA = (
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<dimension ref="A1:B2"/><sheetData>'
    '<row r="1"><c r="A1" t="inlineStr"><is><t>合计</t></is></c>'
    '<c r="B1" t="n"><f>SUM(B2:B2)</f><v>3130</v></c></row>'
    '<row r="2"><c r="A2" t="inlineStr"><is><t>甲</t></is></c>'
    '<c r="B2" t="n"><v>7</v></c></row>'
    "</sheetData></worksheet>"
)

#: 同上，但那一格的**缓存值是长文本**（`t="str"` 是 xlsx 里"公式算出的是字符串"的写法）：
#: 公式原文 `=A1`只有 3 个字符，缓存值 16 个汉字。两遍的正文长度因此**反过来**。
_SHEET_WITH_LONG_CACHED_FORMULA = (
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<dimension ref="A1:B2"/><sheetData>'
    '<row r="1"><c r="A1" t="inlineStr"><is><t>合计</t></is></c>'
    '<c r="B1" t="str"><f>A1</f><v>一二三四五六七八九十一二三四五六</v></c></row>'
    '<row r="2"><c r="A2" t="inlineStr"><is><t>甲</t></is></c>'
    '<c r="B2" t="n"><v>7</v></c></row>'
    "</sheetData></worksheet>"
)


def _xlsx_with_hand_written_sheet(path: Any, sheet_xml: str) -> str:
    """外壳由 openpyxl 生成、`sheet1.xml` 换成手写的那一个 xlsx。

    外壳走库、只换掉那一个成员，是因为 xlsx 的目录结构（`[Content_Types].xml`、
    rels、workbook.xml）自己拼容易拼错，而我们要验的是**那一格**的读法。
    """
    skeleton = io.BytesIO()
    _new_sheet([["甲", 7]]).save(skeleton)
    path.write_bytes(_replace_member(skeleton.getvalue(), "xl/worksheets/sheet1.xml",
                                     sheet_xml.encode("utf-8")))
    return str(path)


def _cached_formula_xlsx(path: Any) -> str:
    """缓存值**比公式原文短**的那一个（`=SUM(B2:B2)` → `3130`）。"""
    return _xlsx_with_hand_written_sheet(path, _SHEET_WITH_CACHED_FORMULA)


def _cached_long_formula_xlsx(path: Any) -> str:
    """缓存值**比公式原文长**的那一个（`=A1` → 16 个汉字）。

    这一支是"两遍长度反过来"的形态，也是 `_Text.reset()` 那处唯一拦得住的形态——
    另一支（缓存值更短）两遍都装得下，换不换对象都绿。
    """
    return _xlsx_with_hand_written_sheet(path, _SHEET_WITH_LONG_CACHED_FORMULA)


def _replace_member(raw: bytes, member: str, data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, data if item.filename == member else source.read(item.filename))
    return out.getvalue()


def _corrupt_member(path: Any, raw: bytes, member: str) -> str:
    """把一个真文件的某个成员换成坏 XML——测"损坏的内部 XML"那一档。"""
    path.write_bytes(_replace_member(raw, member, b"<broken"))
    return str(path)


def _real_bytes(kind: str) -> bytes:
    """三类 OOXML 里任意一类的**好文件字节**（内容非空），用来做"换成坏成员"的底料。

    ⚠ pptx 那一支**必须真的放一张幻灯片**：换了成员却没命中（比如底料是空的
    `Presentation()`，根本没有 `ppt/slides/slide1.xml` 这个成员），"替换"就成了空操作，
    而测试会照样绿——`2026-10-15` 就是这么得出过一个错误结论的。所以调用方那边
    还有一条"底料本身是好的"的断言。
    """
    buffer = io.BytesIO()
    if kind == "docx":
        document = _new_docx()
        document.add_paragraph("底料")
        document.save(buffer)
    elif kind == "xlsx":
        _new_sheet([["底料"]]).save(buffer)
    else:
        prs = _new_pptx()
        _textbox(_slide(prs), "底料")
        prs.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def extractor() -> OfficeExtractor:
    return OfficeExtractor(ExtractConfig())


# ==============================================================================
# 1. 认领与装配
# ==============================================================================


def test_the_four_extensions_are_claimed(extractor):
    missing = sorted(ext for ext in OFFICE_EXTS if not extractor.supports(ext, ""))
    assert missing == []
    assert OFFICE_EXTS == {"pdf", "docx", "xlsx", "pptx"}


def test_every_claimed_extension_has_a_parser(extractor):
    """名单与派发表**必须一致**：少一个 `_PARSERS` 条目，那个格式会走
    `extract()` 里那条"supports() 本该拦住它"的 `UnsupportedFormat`——认领了却读不了，
    而且只在真遇到那个格式时才发现。"""
    assert set(_PARSERS) == set(OFFICE_EXTS)


def test_the_legacy_binary_formats_are_not_claimed(extractor):
    """`.doc` / `.xls` / `.ppt` **不认**：不是坏文件，是 v1 不做的格式。

    认了再抛 `ParseFailed`，等于把"我们没做"记成"用户的文件坏了"——`P4` 的失败报表
    会因此塞满噪音。
    """
    for ext in ("doc", "xls", "ppt", "rtf", "odt", "ods", "odp"):
        assert extractor.supports(ext, "") is False


def test_the_other_extractors_keep_their_extensions(extractor):
    """`plaintext` 与将来的 `web_mail` / `media_meta` 的名单一条都不许抢。

    无扩展名那条尤其要紧：`README` / `Makefile` 归 `plaintext`，被 `office` 抢走
    就会拿去当 PDF 解析。
    """
    for ext in ("txt", "md", "py", "json", "html", "htm", "eml", "msg",
                "jpg", "png", "mp3", "zip", ""):
        assert extractor.supports(ext, "") is False, ext


def test_the_mime_is_never_used(extractor):
    assert extractor.supports("pdf", "text/plain") is True
    assert extractor.supports("txt", "application/pdf") is False


def test_build_follows_the_switch():
    assert build(ExtractConfig(office=True)) is not None
    assert build(ExtractConfig(office=False)) is None


def test_extract_refuses_an_extension_it_never_claimed(tmp_path, extractor):
    """绕过 `supports()` 直接调 `extract()`：**照实说**，不许返回空正文。

    "注册了但返回空"是最难查的一类失败——下游分不出它和"这个文件本来就是空的"。
    """
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 8)     # OLE2 的魔数

    with pytest.raises(UnsupportedFormat) as excinfo:
        extractor.extract(str(path))

    assert "doc" in str(excinfo.value)


def test_the_registry_routes_by_extension(tmp_path):
    """接进注册表之后的整体行为：`.docx` 归 `office`，`.doc` 谁也不认。"""
    registry = build_registry(ExtractConfig())
    docx_path = _docx(tmp_path / "a.docx", _new_docx())
    old_path = tmp_path / "old.doc"
    old_path.write_bytes(b"\xd0\xcf\x11\xe0")

    assert extract_file(docx_path, registry=registry, config=ExtractConfig()).source == "office"
    with pytest.raises(UnsupportedFormat):
        extract_file(str(old_path), registry=registry, config=ExtractConfig())


# ==============================================================================
# 2. 与库自己的标签表对齐
# ==============================================================================


def test_the_hardcoded_word_tags_match_the_library():
    """手写的 Clark 记法必须与 `docx` 的 `qn()` 一字不差。

    写错一个标签名**不会报错**：`iter()` 匹配不上就是收不到那几个字，正文**安静地**
    少一块（比如所有制表符都不见了）。所以这条测试是那几个常量的唯一凭据。
    """
    from docx.oxml.ns import qn

    for own, name in (
        (_W_P, "w:p"),
        (_W_T, "w:t"),
        (_W_TAB, "w:tab"),
        (_W_BR, "w:br"),
        (_W_CR, "w:cr"),
    ):
        assert own == qn(name), name
    assert _W_P.startswith("{" + _W_NS + "}")


# ==============================================================================
# 3. PDF
# ==============================================================================


def test_pages_come_out_in_order_with_markers(tmp_path, extractor):
    path = _pdf(tmp_path / "a.pdf", ["第一页的话", "第二页的话"])

    doc = extractor.extract(path)

    assert doc.text == "【第 1 页】\n第一页的话\n【第 2 页】\n第二页的话"
    assert doc.source == "office"
    assert doc.modality == "text"
    assert doc.warnings == ()
    assert doc.meta["pages"] == 2
    assert doc.meta["empty_pages"] == 0
    assert doc.meta["format"] == "pdf"


def test_the_pdf_title_comes_from_the_metadata(tmp_path, extractor):
    path = _pdf(
        tmp_path / "a.pdf",
        ["正文"],
        metadata={"title": "  年度   报告  ", "producer": "LibreOffice"},
    )

    result = extractor.extract(path)

    # 内部空白压平：标题会进嵌入前缀 `title: {t} | text: ...`（M3），
    # 里面有换行就会与那个格式的分隔符撞上
    assert result.title == "年度 报告"
    assert result.meta["producer"] == "LibreOffice"


def test_a_pdf_without_a_title_says_none(tmp_path, extractor):
    """空标题是**没有标题**，不是"标题是空串"——`title: none` 那个约定（M3）靠它。"""
    path = _pdf(tmp_path / "a.pdf", ["正文"])

    assert extractor.extract(path).title is None


def test_a_blank_page_inside_a_pdf_is_reported(tmp_path, extractor):
    """中间有一页没文字层：**说出来**，而且页数要如实。

    不说的话，症状是"这份 10 页的 PDF 怎么只有 9 页的内容"，而没有任何线索。
    """
    path = _pdf(tmp_path / "a.pdf", ["有字", None, "也有字"])

    doc = extractor.extract(path)

    assert doc.meta["pages"] == 3
    assert doc.meta["empty_pages"] == 1
    assert doc.text == "【第 1 页】\n有字\n【第 3 页】\n也有字"   # 空页连标记都不加
    assert len(doc.warnings) == 1
    assert "1 / 3 页没有文本层" in doc.warnings[0]


def test_a_scanned_pdf_is_an_empty_doc_with_a_warning(tmp_path, extractor):
    """整份都是扫描页：**解析成功、正文为空 + 一条警告**（不是 `ParseFailed`）。

    文件是好的、库也没报错，只是没有文字层。这条与 `plaintext` 的 0 字节文件不同——
    那里"空"是文件自己的事实，这里"空"会让用户以为全文索引是完整的。
    """
    path = _pdf(tmp_path / "scan.pdf", [None, None])

    doc = extractor.extract(path)

    assert doc.text == ""
    assert doc.is_empty is True
    assert doc.truncated is False
    assert doc.meta["empty_pages"] == 2
    assert len(doc.warnings) == 1
    assert "OCR" in doc.warnings[0]


def test_an_encrypted_pdf_is_a_parse_failure(tmp_path, extractor):
    path = _encrypted_pdf(tmp_path / "locked.pdf")

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(path)

    assert "密码" in str(excinfo.value)


def test_a_broken_pdf_is_a_parse_failure(tmp_path, extractor):
    """垃圾字节：`ParseFailed`，而且**消息里带上库报的话**——那是用户手里唯一的线索。"""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"junk" * 50)

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(str(path))

    assert "FileDataError" in str(excinfo.value)


def test_a_pdf_header_with_nothing_behind_it_is_a_parse_failure(tmp_path, extractor):
    """头对、内容空：MuPDF **会照收**，页数是 0（`2026-10-15` 实测）。

    放它过去的话，落点会是"解析成功、正文空 + 一条'每一页都没有文本层'的警告"——
    而它一页都没有，那句警告是**假话**。
    """
    path = tmp_path / "hollow.pdf"
    path.write_bytes(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n")

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(str(path))

    assert "一页都没有" in str(excinfo.value)


def test_an_empty_file_is_a_parse_failure_for_office(tmp_path, extractor):
    """0 字节的 `.pdf` 是**坏文件**，而 0 字节的 `.txt` 是**合法空文档**。

    两边的口径不同是有理由的：容器的头是格式的一部分，空了就什么都没了；而一个文本
    文件本来就可以是空的（`plaintext` 那条测试钉着它的反面）。
    """
    path = tmp_path / "empty.pdf"
    path.write_bytes(b"")

    with pytest.raises(ParseFailed):
        extractor.extract(str(path))


# ==============================================================================
# 4. docx
# ==============================================================================


def test_paragraphs_and_tables_come_out_in_document_order(tmp_path, extractor):
    document = _new_docx()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "格甲"
    table.cell(0, 1).text = "格乙"
    document.add_paragraph("第二段")

    doc = extractor.extract(_docx(tmp_path / "a.docx", document))

    assert doc.text == "第一段\n格甲\n格乙\n第二段"
    assert doc.meta["tables"] == 1


def test_an_empty_paragraph_keeps_its_line(tmp_path, extractor):
    """空段落留一个空行——与 `plaintext` 不清理空行同一个口径。"""
    document = _new_docx()
    document.add_paragraph("一")
    document.add_paragraph("")
    document.add_paragraph("三")

    assert extractor.extract(_docx(tmp_path / "a.docx", document)).text == "一\n\n三"


def test_a_text_box_is_collected_once(tmp_path, extractor):
    """文本框里的字要收，而且**只收一次**。

    文本框的 `w:p` 套在外层段落里（`w:pict > … > `w:txbxContent > w:p`）。外层用
    `iter()` 递归取 `w:t` 时会连带取走里层的字，所以里层那一遍必须跳过——
    不做这一步的症状是同一句话在正文里出现两次（`2026-10-15` 实测踩到过）。
    """
    document = _new_docx()
    document.add_paragraph("正文一段")
    _add_textbox_paragraph(document, "框里的字")

    doc = extractor.extract(_docx(tmp_path / "a.docx", document))

    assert doc.text.count("框里的字") == 1
    assert doc.text == "正文一段\n框里的字"


def test_tabs_and_line_breaks_inside_a_paragraph(tmp_path, extractor):
    """段内的 `w:tab` → `\\t`、`w:br` → `\\n`。不认的话这两种都在正文里消失。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    document = _new_docx()
    paragraph = document.add_paragraph()
    paragraph.add_run("前")
    paragraph.add_run()._r.append(parse_xml(f'<w:tab xmlns:w="{_W_NS}"/>'))
    paragraph.add_run("中")
    paragraph.add_run()._r.append(parse_xml(f'<w:br xmlns:w="{_W_NS}"/>'))
    paragraph.add_run("后")
    assert paragraph.runs                     # 顺带确认 XML 真挂上去了

    assert extractor.extract(_docx(tmp_path / "a.docx", document)).text == "前\t中\n后"


def test_the_docx_title_is_normalized(tmp_path, extractor):
    document = _new_docx()
    document.add_paragraph("正文")
    document.core_properties.title = "  我的   标题  "

    assert extractor.extract(_docx(tmp_path / "a.docx", document)).title == "我的 标题"


def test_a_docx_without_text_is_an_empty_doc_with_a_warning(tmp_path, extractor):
    """只有图片的 docx：解析成功、正文空 + 一条警告（与扫描 PDF 同一档）。"""
    document = _new_docx()

    doc = extractor.extract(_docx(tmp_path / "blank.docx", document))

    assert doc.text == ""
    assert len(doc.warnings) == 1
    assert "OCR" in doc.warnings[0]


# ==============================================================================
# 5. xlsx
# ==============================================================================


def test_sheets_come_out_with_markers_and_empty_cells_are_dropped(tmp_path, extractor):
    workbook = _new_sheet([["项目", "金额"], ["甲", 100], ["合计", "=SUM(B2:B2)"]])

    doc = extractor.extract(_xlsx(tmp_path / "a.xlsx", workbook))

    assert doc.text == (
        "【工作表：台账】\n"
        "项目\t金额\n"
        "甲\t100\n"
        "合计\t=SUM(B2:B2)"
    )


def test_a_hidden_sheet_is_still_collected(tmp_path, extractor):
    """隐藏表照收：数据藏在隐藏表里是常态，"隐藏"是排版决定。"""
    workbook = _new_sheet([["明面上的"]])
    hidden = workbook.create_sheet("隐藏")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "藏起来的字"

    doc = extractor.extract(_xlsx(tmp_path / "a.xlsx", workbook))

    assert doc.meta["sheets"] == 2
    assert "藏起来的字" in doc.text
    assert "【工作表：隐藏】" in doc.text


def test_an_empty_row_is_skipped(tmp_path, extractor):
    """一整行都没有内容的行不产生空行（表格里的空行不是结构，是没填）。"""
    workbook = _new_sheet()
    sheet = workbook.active
    sheet["A1"] = "甲"
    sheet["A3"] = "丙"                      # 第 2 行整行空着

    doc = extractor.extract(_xlsx(tmp_path / "a.xlsx", workbook))

    assert doc.text == "【工作表：台账】\n甲\n丙"


def test_a_workbook_without_formulas_reads_once(tmp_path, extractor):
    """没有公式：第一遍的正文直接用（`formulas` 为 0、没有那条警告）。"""
    doc = extractor.extract(_xlsx(tmp_path / "a.xlsx", _new_sheet([["甲", 1]])))

    assert doc.meta["formulas"] == 0
    assert doc.warnings == ()


def test_a_formula_without_a_cached_value_falls_back_to_the_formula_text(tmp_path, extractor):
    """有公式但文件里没有算好的值（openpyxl / pandas 写的文件就是这样）：

    **退回公式原文 + 一条警告**。让那一格**静默变空**才是真正的坏结果——一整列
    算出来的数就没了，而搜不到不会有任何提示。
    """
    workbook = _new_sheet([["合计", "=SUM(B2:B2)"]])
    path = _xlsx(tmp_path / "a.xlsx", workbook)

    doc = extractor.extract(path)

    assert "=SUM(B2:B2)" in doc.text
    assert doc.meta["formulas"] == 1
    assert len(doc.warnings) == 1
    assert "公式原文" in doc.warnings[0]


def test_a_formula_with_a_cached_value_uses_the_value(tmp_path, extractor):
    """真实 Excel 存出来的文件：**用算好的值**，公式本身不进正文。

    这条同时是"同一个 `read_only` 工作表可以重新迭代"那条载重假设的回归线：第二遍
    要是空转（`2026-10-15` 实测过 openpyxl 3.1.5 支持重迭代），正文里就不会有 3130。
    """
    doc = extractor.extract(_cached_formula_xlsx(tmp_path / "a.xlsx"))

    assert doc.text == "【工作表：台账】\n合计\t3130\n甲\t7"
    assert "=SUM" not in doc.text
    assert doc.meta["formulas"] == 1        # 第一遍数出来的公式格数
    assert doc.warnings == ()               # 缓存值拿到了，就没有"退回原文"那条警告


def test_the_xlsx_title_comes_from_the_properties(tmp_path, extractor):
    workbook = _new_sheet([["甲"]])
    workbook.properties.title = "台账本"

    assert extractor.extract(_xlsx(tmp_path / "a.xlsx", workbook)).title == "台账本"


# ==============================================================================
# 6. pptx
# ==============================================================================


def test_slides_come_out_with_markers_in_order(tmp_path, extractor):
    prs = _new_pptx()
    _textbox(_slide(prs), "第一张的话")
    _textbox(_slide(prs), "第二张的话")

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert doc.text == "【第 1 张幻灯片】\n第一张的话\n【第 2 张幻灯片】\n第二张的话"
    assert doc.meta["slides"] == 2


def test_a_table_becomes_one_line_per_row_with_tabs(tmp_path, extractor):
    prs = _new_pptx()
    slide = _slide(prs)
    table = slide.shapes.add_table(2, 2, 0, 0, 300, 100).table
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "甲"
    table.cell(1, 1).text = ""              # 空格子丢掉，不用制表符占位

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert doc.text == "【第 1 张幻灯片】\n项目\t金额\n甲"


def test_a_group_shape_is_walked_into(tmp_path, extractor):
    """组合形状里的字要收——只走顶层 `slide.shapes` 的话它整块都不见了。"""
    prs = _new_pptx()
    slide = _slide(prs)
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, 100, 50).text_frame.text = "组里的字"

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert doc.text == "【第 1 张幻灯片】\n组里的字"


def test_speaker_notes_are_collected_with_a_marker(tmp_path, extractor):
    """演讲者备注要收：那是用户真的写过、也真的会去搜的字。"""
    prs = _new_pptx()
    slide = _slide(prs)
    _textbox(slide, "正文")
    slide.notes_slide.notes_text_frame.text = "备注里的话"

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert doc.text == "【第 1 张幻灯片】\n正文\n【第 1 张幻灯片·备注】\n备注里的话"
    assert doc.meta["notes"] == 1


def test_a_slide_without_notes_gets_no_marker(tmp_path, extractor):
    """没有备注就不加标记——空标记行是噪音，不是信息。"""
    prs = _new_pptx()
    _textbox(_slide(prs), "只有正文")

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert "备注" not in doc.text
    assert doc.meta["notes"] == 0


def test_an_image_only_pptx_is_an_empty_doc_with_a_warning(tmp_path, extractor):
    """只有图片的 pptx：标记照加，但**那条空正文警告要在**。

    ⚠ `doc.is_empty` 在这里是 `False`（`text` 里有标记行，而标记行也是文字），所以
    "有没有正文"这件事**不能**用 `is_empty` 判——它由 `_Text.has_body` 回答。第一版
    实现就是拿 `text.strip()` 判的，标记把它骗过去了，这条测试当场红。
    """
    prs = _new_pptx()
    _slide(prs)

    doc = extractor.extract(_save_pptx(prs, tmp_path / "a.pptx"))

    assert doc.text == "【第 1 张幻灯片】"     # 标记还是加，只是没内容
    assert len(doc.warnings) == 1
    assert "没有提取到任何文本" in doc.warnings[0]


def test_a_marker_alone_is_not_body_text(tmp_path, extractor):
    """同样的道理在 xlsx 上：一个空工作簿只剩工作表标记，警告不能因此消失。"""
    doc = extractor.extract(_xlsx(tmp_path / "empty.xlsx", _new_sheet()))

    assert doc.text == "【工作表：台账】"
    assert len(doc.warnings) == 1
    assert "没有提取到任何文本" in doc.warnings[0]


# ==============================================================================
# 7. 上限与截断
# ==============================================================================


def test_the_cap_stops_at_a_line_boundary(tmp_path):
    """到上限**整行地停**：`truncated` 为真、留一条警告、正文不会超过上限。"""
    workbook = _new_sheet([["第{}行".format(n)] for n in range(1, 21)])
    path = _xlsx(tmp_path / "a.xlsx", workbook)
    capped = OfficeExtractor(ExtractConfig(), max_chars=30)

    doc = capped.extract(path)

    assert doc.truncated is True
    assert len(doc.text) <= 30
    assert doc.text.startswith("【工作表：台账】")   # 从第一行开始，不是被砍掉的头
    assert len(doc.warnings) == 1
    assert "上限" in doc.warnings[0]
    # 截断是"没提取完"，不是"文件坏了"：前面那些行照旧好用
    assert "第1行" in doc.text


def test_the_truncation_flag_describes_the_second_pass(tmp_path):
    """`truncated` 说的是**最终那份正文**（第二遍），不是第一遍那份。

    这个工作簿的两遍长度**反过来**：第一遍是公式原文（`合计\\t=A1`，共 19 字符），
    第二遍是缓存值（16 个汉字，共 28 字符）。上限取 30 时第一遍装得下、第二遍装不下，
    正文停在第 2 行——`truncated` 与那条警告**都必须跟着第二遍**。

    第一版实现里 xlsx 的第二遍是**另造一个 `_Text`**，而 `extract()` 读的是第一个对象，
    于是它给出的是"正文只到第 2 行、`truncated=False`、`warnings=()`"（`2026-10-15`
    实测）。这条测试就是 `_Text.reset()` 的回归线。

    ⚠ 拦得住这一支的**只有"缓存值更长"**：第一版这条测试用的是缓存值**更短**的工作簿
    （`=SUM(B2:B2)` → `3130`），两遍都装得下，换不换对象都绿。写测试时按"更短"算的
    那笔账也是错的——那个工作簿的第一遍是 27 字符，上限 28 **装得下**，测试从写下
    那一刻起就永远绿（`2026-10-15` 复查时把旧写法临时换回去才发现）。
    """
    path = _cached_long_formula_xlsx(tmp_path / "a.xlsx")

    doc = OfficeExtractor(ExtractConfig(), max_chars=30).extract(path)

    assert doc.text == "【工作表：台账】\n合计\t一二三四五六七八九十一二三四五六"
    assert "甲" not in doc.text                   # 第 3 行确实没进来，不是自说自话
    assert doc.truncated is True                  # 第二遍装不下，就要说装不下
    assert len(doc.warnings) == 1
    assert "上限" in doc.warnings[0]


def test_a_full_first_pass_skips_the_second_pass(tmp_path):
    """第一遍就到顶时**不读第二遍**：正文留着公式原文，`truncated` 照实说满。

    这条不只是省一次读——`_Text.reset()` 会**清掉** `truncated`，所以"第二遍只在第一遍
    没满时才跑"这道闸门是那句 `reset()` 的前提。闸门一旦被撤掉，一份满了的正文会被
    第二遍重置成"没满"，`truncated` 与警告一起消失（`2026-10-15` 实测：上限 25 时正文
    仍停在第 2 行，标志变成假、警告也没了）。
    """
    path = _cached_formula_xlsx(tmp_path / "a.xlsx")

    doc = OfficeExtractor(ExtractConfig(), max_chars=25).extract(path)

    assert doc.text == "【工作表：台账】\n合计\t=SUM(B2:B2)"
    assert doc.meta["formulas"] == 1
    assert "3130" not in doc.text                 # 第二遍没跑，缓存值没进来
    assert doc.truncated is True
    assert len(doc.warnings) == 1
    assert "上限" in doc.warnings[0]


def test_the_cap_does_not_fire_on_a_normal_file(tmp_path, extractor):
    """正常文件不许有 `truncated` / 警告——默认上限是 1e6 字符，
    而 `D8` 的 200 片上限最多只会用到它的三成（算术在模块文档里）。"""
    doc = extractor.extract(_xlsx(tmp_path / "a.xlsx", _new_sheet([["甲", 1]])))

    assert doc.truncated is False
    assert doc.warnings == ()
    assert MAX_CHARS == 1_000_000


def test_the_file_size_limit_also_applies_to_office(tmp_path):
    """`[scan] max_file_mb` 那条上限由 `read_bytes()` 管，四类提取器共用——
    office 不另开一个口子（上限只有一个真相来源）。"""
    path = _pdf(tmp_path / "big.pdf", ["正文"])
    tiny = OfficeExtractor(ExtractConfig(max_bytes=10))

    with pytest.raises(FileUnreadable):
        tiny.extract(path)


# ==============================================================================
# 8. 坏文件：五种形态，一个都不许把整轮索引带下去
# ==============================================================================


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_a_file_that_is_not_a_zip_is_a_parse_failure(tmp_path, extractor, kind):
    path = tmp_path / f"broken.{kind}"
    path.write_bytes(b"this is not a zip file at all" * 4)

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(str(path))

    assert "BadZipFile" in str(excinfo.value)


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_a_zip_that_is_not_ooxml_is_a_parse_failure(tmp_path, extractor, kind):
    """是合法 zip，但没有 `[Content_Types].xml`：库抛的是 `KeyError`。

    一个 `.docx` 实际上是别的东西（改名来的 zip）在盘上很常见，这条要接住。
    """
    path = tmp_path / f"fake.{kind}"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "我不是 OOXML")
    path.write_bytes(buffer.getvalue())

    with pytest.raises(ParseFailed):
        extractor.extract(str(path))


def test_a_broken_document_xml_is_a_parse_failure(tmp_path, extractor):
    """坏 `document.xml`：python-docx 走 lxml，抛的是 `XMLSyntaxError`。"""
    path = _corrupt_member(tmp_path / "c.docx", _real_bytes("docx"), "word/document.xml")

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(path)

    assert "XMLSyntaxError" in str(excinfo.value)


def test_a_broken_sheet_xml_is_a_parse_failure(tmp_path, extractor):
    """坏 `sheet1.xml`：openpyxl 走 stdlib 的 `ElementTree`，抛的是 `ParseError`。

    与上一条分开列，是因为**两族的异常类没有共同的父亲**（lxml 的
    `XMLSyntaxError` 不是 `ET.ParseError` 的子类，实测过 mro）。少接一个的后果是那种
    坏文件变成未捕获异常，把整轮索引带下去。
    """
    path = _corrupt_member(
        tmp_path / "c.xlsx", _real_bytes("xlsx"), "xl/worksheets/sheet1.xml"
    )

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(path)

    assert "ParseError" in str(excinfo.value)


def test_a_broken_slide_xml_is_a_parse_failure(tmp_path, extractor):
    """坏 `slide1.xml`：python-pptx 与 python-docx 同族，抛 `XMLSyntaxError`。

    ⚠ 第一条断言是**底料的凭据**，不是多余的：这张 pptx 必须真有 `ppt/slides/slide1.xml`
    这个成员，"换成坏 XML"才命中。用空的 `Presentation()` 当底料的话，替换是空操作、
    测试照样绿——`2026-10-15` 就是这么得出过一个错误结论（"python-pptx 会静默丢掉
    坏幻灯片"），重做时加上这条断言才发现它其实照抛。
    """
    good = tmp_path / "good.pptx"
    good.write_bytes(_real_bytes("pptx"))
    assert extractor.extract(str(good)).meta["slides"] == 1     # 底料本身是好的

    path = _corrupt_member(tmp_path / "c.pptx", _real_bytes("pptx"), "ppt/slides/slide1.xml")

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(path)

    assert "XMLSyntaxError" in str(excinfo.value)


def test_the_failure_reason_is_always_something_a_human_can_read(tmp_path, extractor):
    """消息里要有**库的名字与它自己说的话**——"解析失败"四个字等于没写
    （`files.error_msg` 是用户唯一能看见的东西）。"""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"junk" * 50)

    with pytest.raises(ParseFailed) as excinfo:
        extractor.extract(str(path))

    message = str(excinfo.value)
    assert "PDF" in message
    assert "FileDataError" in message


# ==============================================================================
# 9. 失败的分类与结果的形状
# ==============================================================================


def test_parse_failed_is_not_unsupported_format():
    """两者必须是**两个类**（`P4` 的处置相反）：`ParseFailed` 进失败报表，
    `UnsupportedFormat` 留在 `NAME_ONLY`。合成一个就同时坏两头。"""
    assert issubclass(ParseFailed, ExtractError)
    assert issubclass(UnsupportedFormat, ExtractError)
    assert not issubclass(ParseFailed, UnsupportedFormat)
    assert not issubclass(UnsupportedFormat, ParseFailed)


def test_the_result_is_a_labelled_frozen_extracted_doc(tmp_path, extractor):
    import dataclasses

    doc = extractor.extract(_docx(tmp_path / "a.docx", _new_docx()))

    assert isinstance(doc, ExtractedDoc)
    assert doc.source == "office"
    assert doc.modality == "text"          # v1 全是 text（core/models.py 的口径）
    assert doc.meta["bytes"] > 0
    assert doc.meta["chars"] == len(doc.text)
    with pytest.raises(dataclasses.FrozenInstanceError):
        doc.text = "别的"
